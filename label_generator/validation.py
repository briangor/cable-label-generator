import math
from copy import deepcopy

from docx import Document

from .extraction import CableRecord

DEFAULT_TOP_LABEL_STARTS = (1, 4, 7, 10, 13)
DEFAULT_BOTTOM_LABEL_STARTS = (0, 3, 6, 9, 12)
DEFAULT_LABEL_ROWS = (0, 2, 4, 6, 8, 10)
DEFAULT_LABELS_PER_TABLE = 30

def duplicate_label_groups(
    records: list[CableRecord],
    group_size: int = 5,
) -> list[CableRecord]:
    """
    Duplicate each physical label group for the two cable ends.

    ``group_size`` describes the physical template arrangement. It does not
    change the invariant that every cable record must produce exactly two
    physical labels.
    """
    if group_size <= 0:
        raise ValueError("group_size must be greater than zero.")

    stream: list[CableRecord] = []

    for start in range(0, len(records), group_size):
        chunk = records[start:start + group_size]
        stream.extend(chunk)
        stream.extend(chunk)

    return stream


def validate_two_labels_per_cable(
    records: list[CableRecord],
    label_stream: list[CableRecord],
) -> None:
    """Validate the two-physical-label invariant for every cable."""
    expected = len(records) * 2

    if len(label_stream) != expected:
        raise ValueError(
            "Two-label-per-cable validation failed: "
            f"{len(records)} cables require {expected} physical labels, "
            f"but {len(label_stream)} were generated."
        )

    counts: dict[tuple[str, int], int] = {}

    for record in label_stream:
        key = (record[2], record[0])
        counts[key] = counts.get(key, 0) + 1

    for record in records:
        key = (record[2], record[0])
        count = counts.get(key, 0)

        if count != 2:
            raise ValueError(
                "Two-label-per-cable validation failed for "
                f"{record[2]} / {record[3]}: expected 2, found {count}."
            )

    if len(counts) != len(records):
        raise ValueError(
            "Two-label-per-cable validation failed: the physical label "
            "stream contains records that are not present in the cable set."
        )


# ---------------------------------------------------------------------------
# DOCX rendering
# ---------------------------------------------------------------------------


def layout_config(profile: dict) -> dict:
    """Return validated physical-layout settings from the profile."""
    layout = profile.get("layout", {})

    label_rows = tuple(
        int(value)
        for value in layout.get("label_rows", DEFAULT_LABEL_ROWS)
    )
    top_starts = tuple(
        int(value)
        for value in layout.get(
            "top_label_starts",
            DEFAULT_TOP_LABEL_STARTS,
        )
    )
    bottom_starts = tuple(
        int(value)
        for value in layout.get(
            "bottom_label_starts",
            DEFAULT_BOTTOM_LABEL_STARTS,
        )
    )
    labels_per_table = int(
        layout.get("labels_per_table", DEFAULT_LABELS_PER_TABLE)
    )

    if not label_rows:
        raise ValueError("layout.label_rows must not be empty.")

    if not top_starts or not bottom_starts:
        raise ValueError(
            "layout.top_label_starts and bottom_label_starts "
            "must not be empty."
        )

    if len(top_starts) != len(bottom_starts):
        raise ValueError(
            "Top and bottom label rows must contain the same "
            "number of label positions."
        )

    calculated = len(label_rows) * len(top_starts)

    if labels_per_table != calculated:
        raise ValueError(
            "layout.labels_per_table does not match the configured "
            f"label grid: expected {calculated}, got {labels_per_table}."
        )

    return {
        "label_rows": label_rows,
        "top_label_starts": top_starts,
        "bottom_label_starts": bottom_starts,
        "labels_per_table": labels_per_table,
        "pages_per_table": int(layout.get("pages_per_table", 2)),
    }


def _layout_table_xml(table):
    """Return table XML with cell content removed for geometry comparison."""
    xml = deepcopy(table._tbl)
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    for cell in xml.iter(f"{{{namespace}}}tc"):
        for child in list(cell):
            if child.tag != f"{{{namespace}}}tcPr":
                cell.remove(child)

    return xml


def validate_template_layout(
    template_doc: Document,
    profile: dict,
) -> None:
    """Validate that the template can satisfy the configured label grid."""
    layout = layout_config(profile)

    if not template_doc.tables:
        raise ValueError("The template DOCX does not contain a table.")

    master = template_doc.tables[0]

    expected_columns = (
        max(
            max(layout["top_label_starts"]),
            max(layout["bottom_label_starts"]),
        )
        + 2
    )

    if len(master.rows) <= max(layout["label_rows"]):
        raise ValueError(
            "Template table does not contain all configured label rows."
        )

    if len(master.columns) < expected_columns:
        raise ValueError(
            "Template table does not contain enough columns for the "
            "configured label positions."
        )

    if layout["pages_per_table"] <= 0:
        raise ValueError(
            "layout.pages_per_table must be greater than zero."
        )


def validate_generated_layout(
    template_doc: Document,
    generated_doc: Document,
    profile: dict,
    label_count: int,
) -> None:
    """Validate generated structure and physical table geometry."""
    layout = layout_config(profile)
    labels_per_table = layout["labels_per_table"]
    expected_tables = math.ceil(label_count / labels_per_table)

    if len(generated_doc.tables) != expected_tables:
        raise ValueError(
            "Generated table count does not match label capacity: "
            f"expected {expected_tables}, got {len(generated_doc.tables)}."
        )

    page_breaks = sum(
        1
        for paragraph in generated_doc.paragraphs
        if paragraph._p.xpath(".//w:br[@w:type='page']")
    )

    expected_page_breaks = max(0, expected_tables - 1)

    if page_breaks != expected_page_breaks:
        raise ValueError(
            "Generated page-break count is incorrect: "
            f"expected {expected_page_breaks}, got {page_breaks}."
        )

    master_xml = _layout_table_xml(template_doc.tables[0])

    for index, table in enumerate(generated_doc.tables, start=1):
        if _layout_table_xml(table).xml != master_xml.xml:
            raise ValueError(
                f"Generated table {index} does not preserve the "
                "template's physical table geometry."
            )

    template_section = template_doc.sections[0]._sectPr
    generated_section = generated_doc.sections[0]._sectPr

    if template_section.xml != generated_section.xml:
        raise ValueError(
            "Generated document section/page settings differ from "
            "the template."
        )


