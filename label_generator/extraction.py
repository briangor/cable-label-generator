import re
from openpyxl.utils.cell import column_index_from_string

CableRecord = tuple[int, str, str, str]

def parse_port_number(value) -> int | None:
    """Extract the numeric port suffix from values such as 25GE1/0/17."""
    if value is None:
        return None

    match = re.search(r"/(\d+)\s*$", str(value).strip())
    return int(match.group(1)) if match else None


def column_number(column: str) -> int:
    """Convert an Excel column letter to a 1-based column number."""
    try:
        return column_index_from_string(column)
    except ValueError as exc:
        raise ValueError(f"Invalid Excel column: {column!r}") from exc


def cell_value(ws, row: int, column: str):
    """Read a worksheet cell using an Excel column letter."""
    return ws.cell(row=row, column=column_number(column)).value


def find_section_starts(ws, section_config: dict) -> list[int]:
    """Find section header rows using profile-defined rules."""
    device_column = section_config["device_column"]
    header_row_offset = int(section_config.get("header_row_offset", 1))
    header_value = section_config.get("header_value")
    header_match = section_config.get("header_match")
    pattern = re.compile(header_match) if header_match else None

    starts = []

    for row in range(1, ws.max_row + 1 - header_row_offset):
        device = cell_value(ws, row, device_column)
        header = cell_value(ws, row + header_row_offset, device_column)

        if not isinstance(device, str):
            continue

        if pattern and not pattern.search(device):
            continue

        if header_value is not None and header != header_value:
            continue

        starts.append(row)

    return starts


def extract_sections(ws, profile: dict) -> list[dict]:
    """Extract all configured cable sides from all detected sections."""
    section_config = profile["sections"]
    side_configs = profile["sides"]
    starts = find_section_starts(ws, section_config)
    data_start_offset = int(section_config.get("data_start_offset", 2))

    sections = []

    for index, start in enumerate(starts):
        end = starts[index + 1] - 1 if index + 1 < len(starts) else ws.max_row

        sides = []

        for side in side_configs:
            device_column = side["device_column"]
            port_column = side["port_column"]
            label_column = side["label_column"]

            device_value = cell_value(ws, start, device_column)
            device = str(device_value).strip() if device_value is not None else ""

            records: list[CableRecord] = []

            for row in range(start + data_start_offset, end + 1):
                local_port = cell_value(ws, row, port_column)
                peer_label = cell_value(ws, row, label_column)

                if local_port in (None, "") or peer_label in (None, ""):
                    continue

                port_number = parse_port_number(local_port)
                if port_number is None:
                    continue

                records.append(
                    (
                        port_number,
                        str(peer_label).strip(),
                        device,
                        str(local_port).strip(),
                    )
                )

            if device or records:
                sides.append(
                    {
                        "name": side.get("name", "side"),
                        "device": device,
                        "records": records,
                    }
                )

        sections.append({"start": start, "sides": sides})

    return sections


def validate_sheet_structure(ws, profile: dict) -> None:
    """Validate that a selected worksheet matches the profile's structure."""
    section_config = profile.get("sections")
    side_configs = profile.get("sides")

    if not isinstance(section_config, dict):
        raise ValueError("Profile is missing a valid 'sections' configuration.")
    if not isinstance(side_configs, list) or not side_configs:
        raise ValueError("Profile must define at least one cable side.")

    starts = find_section_starts(ws, section_config)
    if not starts:
        raise ValueError("no matching cable sections were found")

    missing = []
    for side in side_configs:
        for key in ("device_column", "port_column", "label_column"):
            if not side.get(key):
                missing.append(f"{side.get('name', 'side')}.{key}")
    if missing:
        raise ValueError("missing side configuration: " + ", ".join(missing))

    for start in starts:
        for side in side_configs:
            header_offset = int(section_config.get("header_row_offset", 1))
            expected = side.get("header_value", section_config.get("header_value"))
            if expected is None:
                continue
            header = cell_value(ws, start + header_offset, side["port_column"])
            if header != expected:
                raise ValueError(
                    f"expected {expected!r} at row {start + header_offset} "
                    f"for {side.get('name', 'side')}, got {header!r}"
                )

    data_start_offset = int(section_config.get("data_start_offset", 2))
    has_data = False
    for start in starts:
        end = ws.max_row
        for row in range(start + data_start_offset, end + 1):
            for side in side_configs:
                local_port = cell_value(ws, row, side["port_column"])
                peer_label = cell_value(ws, row, side["label_column"])
                if local_port not in (None, "") and peer_label not in (None, ""):
                    if parse_port_number(local_port) is not None:
                        has_data = True
                        break
            if has_data:
                break
        if has_data:
            break

    if not has_data:
        raise ValueError("matching sections were found, but no valid cable records were found")

