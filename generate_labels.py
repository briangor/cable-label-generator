#!/usr/bin/env python3
"""Generate printable cable labels from an XLSX workbook and DOCX template."""

from __future__ import annotations

__author__ = "Brian Gor"

import argparse
import math
import re
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from label_generator import (
    duplicate_label_groups,
    extract_sections,
    generate_document,
    layout_config,
    load_profile,
    order_records_by_profile,
    validate_generated_layout,
    validate_sheet_structure,
    validate_two_labels_per_cable,
)


def _confirm(prompt: str) -> bool:
    return input(f"{prompt} [Y/n]: ").strip().lower() in ("", "y", "yes")


def _choose(items: list[str], title: str) -> str:
    if not items:
        raise ValueError(f"No {title.lower()} found.")
    print(f"\n{title}")
    for index, item in enumerate(items, 1):
        print(f"  {index}. {item}")
    while True:
        try:
            index = int(input("Select number: ").strip())
        except ValueError:
            print("Please enter a number.")
            continue
        if 1 <= index <= len(items):
            return items[index - 1]
        print("Please select one of the listed numbers.")


def _sanitize_filename_suffix(value: str | None) -> str | None:
    """Return a safe filename suffix, or None when no suffix is supplied."""
    if value is None:
        return None

    suffix = value.strip()
    if not suffix:
        return None

    suffix = re.sub(r"[^A-Za-z0-9._-]+", "_", suffix)
    suffix = re.sub(r"_+", "_", suffix).strip("._-")
    if not suffix:
        raise ValueError("Filename suffix must contain at least one valid character.")

    return suffix


def _default_output_path(suffix: str | None = None) -> Path:
    """Return the standard generated-label output path."""
    timestamp = datetime.now().strftime("%Y%m%d-%H%M")
    safe_suffix = _sanitize_filename_suffix(suffix)
    filename = f"labels_{timestamp}"
    if safe_suffix:
        filename += f"_{safe_suffix}"
    return Path("labels") / f"{filename}.docx"


def _resolve_output_path(
    output: Path | None,
    suffix: str | None = None,
) -> Path:
    """Use an explicit output path or the standard labels/ timestamp path."""
    if output is not None:
        return output
    return _default_output_path(suffix)


def _discover_workbooks(directory: Path) -> list[Path]:
    return sorted(
        p for p in directory.glob("*.xlsx") if not p.name.startswith("~$")
    )


def _select_sheet(workbook: Path) -> str:
    """Choose a worksheet without assuming which profile describes it."""
    book = load_workbook(workbook, read_only=False, data_only=True)
    sheets = book.sheetnames
    selected = sheets[0] if len(sheets) == 1 else _choose(sheets, "Available worksheets")
    print(f"Selected worksheet: {selected}")
    if not _confirm("Use this worksheet?"):
        raise SystemExit("Generation cancelled.")
    return selected


def _discover_profiles(directory: Path) -> list[Path]:
    """Find YAML profiles available for interactive generation."""
    return sorted(
        [*directory.glob("*.yaml"), *directory.glob("*.yml")],
        key=lambda path: path.name.lower(),
    )


def _select_profile(directory: Path, workbook: Path, sheet: str) -> dict:
    """Choose a profile that can validate the selected worksheet."""
    profile_paths = _discover_profiles(directory)
    if not profile_paths:
        raise ValueError(f"No YAML profiles found in {directory}.")

    book = load_workbook(workbook, read_only=False, data_only=True)
    compatible = []
    for path in profile_paths:
        try:
            candidate = load_profile(path)
            candidate = deepcopy(candidate)
            candidate["workbook"] = deepcopy(candidate["workbook"])
            candidate["workbook"]["sheet"] = sheet
            validate_sheet_structure(book[sheet], candidate)
            label = f"{path.name} ({candidate.get('name', path.stem)})"
            compatible.append((label, candidate))
        except (ValueError, KeyError, TypeError):
            continue

    if not compatible:
        raise ValueError(
            f'No profiles in "{directory}" match worksheet "{sheet}". '
            "Check the profile's section and side configuration."
        )

    labels = [label for label, _ in compatible]
    selected_label = _choose(labels, f"Profiles compatible with worksheet {sheet}")
    profile = next(profile for label, profile in compatible if label == selected_label)
    print(f"Selected profile: {selected_label}")
    print("Sheet structure: valid for selected profile")
    return profile


def _review_inventory(workbook: Path, profile: dict) -> None:
    book = load_workbook(workbook, data_only=True, read_only=False)
    sheet = profile["workbook"]["sheet"]
    validate_sheet_structure(book[sheet], profile)
    sections = extract_sections(book[sheet], profile)
    if not sections:
        raise ValueError(f'No cable sections were found in sheet "{sheet}".')

    inventory = {}
    cables = 0
    for section in sections:
        for side in section["sides"]:
            device = side["device"]
            count = len(side["records"])
            if not device:
                continue
            entry = inventory.setdefault(
                device,
                {"rack": "-", "switch_type": "-", "switch_name": device, "cables": 0},
            )
            entry["cables"] += count
            cables += count

    inventory_config = profile.get("inventory", {})
    device_pattern = inventory_config.get("device_pattern")
    if device_pattern:
        import re
        pattern = re.compile(device_pattern)
        for device, entry in inventory.items():
            match = pattern.search(device)
            if match:
                groups = match.groupdict()
                entry["rack"] = groups.get("rack", entry["rack"])
                entry["switch_type"] = groups.get("switch_type", entry["switch_type"])
                entry["switch_name"] = groups.get("switch_name", entry["switch_name"])

    labels = cables * 2
    layout = layout_config(profile)
    tables = math.ceil(labels / layout["labels_per_table"])
    pages = tables * layout["pages_per_table"]

    print("\nSwitch inventory")
    print("  Rack     Type   Switch      Cables")
    print("  -------  -----  ----------  ------")
    for entry in inventory.values():
        print(
            f"  {entry['rack']:<7}  {entry['switch_type']:<5}  "
            f"{entry['switch_name']:<10}  {entry['cables']:>6}"
        )

    print("\nGeneration summary")
    print(f"  Total cables:          {cables}")
    print(f"  Physical labels:       {labels}")
    print(f"  Expected tables:       {tables}")
    print(f"  Expected pages:        {pages}")


def _interactive(args):
    workbooks = _discover_workbooks(args.workbook_dir)
    workbook = args.workbook_dir / _choose(
        [p.name for p in workbooks],
        f"XLSX workbooks in {args.workbook_dir}",
    )
    if not _confirm(f"Use workbook {workbook.name}?"):
        raise SystemExit("Generation cancelled.")

    selected_sheet = _select_sheet(workbook)
    profile = _select_profile(args.profile_dir, workbook, selected_sheet)
    _review_inventory(workbook, profile)

    if not _confirm("Proceed with generation?"):
        raise SystemExit("Generation cancelled.")

    suffix = args.suffix
    if args.output_docx is None and suffix is None:
        suffix = input(
            "Custom filename suffix (optional, e.g. NBO or KIS): "
        ).strip() or None

    output = _resolve_output_path(args.output_docx, suffix)
    return workbook, output, args.template, profile


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate cable labels from an XLSX workbook."
    )
    parser.add_argument("input_xlsx", nargs="?", type=Path)
    parser.add_argument("output_docx", nargs="?", type=Path)
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("templates/template.docx"),
        help="DOCX used as the formatting/layout template.",
    )
    parser.add_argument(
        "--profile",
        type=Path,
        default=None,
        help=(
            "YAML profile for explicit/non-interactive mode "
            "(default: profiles/example.yaml). Interactive mode prompts for a profile."
        ),
    )
    parser.add_argument(
        "--profile-dir",
        type=Path,
        default=Path("profiles"),
        help="Directory containing YAML profiles for interactive selection (default: profiles/).",
    )
    parser.add_argument(
        "--workbook-dir",
        type=Path,
        default=Path("workbook"),
        help="Directory searched for XLSX files in interactive mode (default: workbook/).",
    )
    parser.add_argument(
        "--suffix",
        help=(
            "Optional suffix appended to the default filename, e.g. "
            "--suffix NBO -> labels_YYYYMMDD-HHMM_NBO.docx."
        ),
    )
    return parser


def main() -> None:
    parser = _build_parser()
    args = parser.parse_args()
    if args.input_xlsx is None:
        input_xlsx, output_docx, template, profile = _interactive(args)
    else:
        profile_path = args.profile or Path("profiles/example.yaml")
        profile = load_profile(profile_path)
        input_xlsx = args.input_xlsx
        output_docx = _resolve_output_path(args.output_docx, args.suffix)
        template = args.template

    cable_count, label_count, change_report = generate_document(
        input_xlsx, output_docx, template, profile
    )

    tables = math.ceil(label_count / layout_config(profile)["labels_per_table"])
    print(f"\nProfile:      {profile.get('name', (args.profile or Path('profiles/example.yaml')).stem)}")
    print(f"Source sheet: {profile['workbook']['sheet']}")
    print(f"Cables used:  {cable_count}")
    print(f"Labels:       {label_count}")
    print(f"Tables:       {tables}")
    print(f"Pages:        {tables * layout_config(profile)['pages_per_table']}")

    if change_report["reference"]:
        print(f"Reference records retained: {len(change_report['retained'])}")
        print(f"Reference records removed:  {len(change_report['removed'])}")
        print(f"New records:                {len(change_report['added'])}")

        if change_report["removed"]:
            print("Removed reference records:")
            for device, port in sorted(change_report["removed"]):
                print(f"  - {device}: {port}")

        if change_report["added"]:
            print("New workbook records:")
            for device, port in sorted(change_report["added"]):
                print(f"  + {device}: {port}")

    print(f"Output:       {output_docx}")


if __name__ == "__main__":
    main()
