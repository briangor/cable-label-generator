from pathlib import Path

from openpyxl import load_workbook

from label_generator import load_profile, validate_sheet_structure

ROOT = Path(__file__).resolve().parents[1]
PROFILE = ROOT / "profiles" / "example.yaml"
WORKBOOK = ROOT / "workbook" / "example.xlsx"


def test_selected_example_sheet_structure_is_valid():
    profile = load_profile(PROFILE)
    workbook = load_workbook(WORKBOOK, data_only=True)
    validate_sheet_structure(workbook[profile["workbook"]["sheet"]], profile)


def test_invalid_sheet_structure_is_rejected(tmp_path):
    profile = load_profile(PROFILE)
    workbook = load_workbook(WORKBOOK)
    ws = workbook.create_sheet("Invalid")
    ws["A1"] = "not a cable section"
    path = tmp_path / "invalid.xlsx"
    workbook.save(path)

    workbook = load_workbook(path, data_only=True)
    try:
        validate_sheet_structure(workbook["Invalid"], profile)
    except ValueError as exc:
        assert "no matching cable sections" in str(exc)
    else:
        raise AssertionError("Invalid worksheet should have been rejected")


def test_declining_final_confirmation_does_not_generate_docx(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import generate_labels

    profile = load_profile(PROFILE)
    args = SimpleNamespace(
        workbook_dir=WORKBOOK.parent,
        profile_dir=ROOT / "profiles",
        output_docx=tmp_path / "labels" / "labels_20261008-1200.docx",
        template=ROOT / "templates" / "template.docx",
        suffix=None,
    )

    confirmations = iter([True, True, False])
    monkeypatch.setattr(generate_labels, "_confirm", lambda prompt: next(confirmations))
    monkeypatch.setattr(generate_labels, "_choose", lambda items, title: items[0])

    generate_calls = []
    monkeypatch.setattr(
        generate_labels,
        "generate_document",
        lambda *args, **kwargs: generate_calls.append((args, kwargs)),
    )

    try:
        generate_labels._interactive(args)
    except SystemExit as exc:
        assert str(exc) == "Generation cancelled."
    else:
        raise AssertionError("Interactive generation should have been cancelled")

    assert generate_calls == []
    assert not args.output_docx.exists()


def test_interactive_custom_suffix_is_used(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import generate_labels

    profile = load_profile(PROFILE)
    args = SimpleNamespace(
        workbook_dir=WORKBOOK.parent,
        profile_dir=ROOT / "profiles",
        output_docx=None,
        suffix=None,
        template=ROOT / "templates" / "template.docx",
    )

    confirmations = iter([True, True, True])
    monkeypatch.setattr(generate_labels, "_confirm", lambda prompt: next(confirmations))
    monkeypatch.setattr(generate_labels, "_choose", lambda items, title: items[0])
    monkeypatch.setattr("builtins.input", lambda prompt: "NBO")

    workbook, output, template, selected_profile = generate_labels._interactive(args)

    assert workbook == WORKBOOK
    assert output.parent == Path("labels")
    assert output.name.endswith("_NBO.docx")
    assert template == args.template
    assert selected_profile["workbook"]["sheet"] == profile["workbook"]["sheet"]
    assert selected_profile["name"] == profile["name"]


def test_profile_discovery_includes_yaml_and_yml(tmp_path):
    import generate_labels

    (tmp_path / "second.yml").write_text("workbook: {}\n", encoding="utf-8")
    (tmp_path / "first.yaml").write_text("workbook: {}\n", encoding="utf-8")
    (tmp_path / "ignored.txt").write_text("not a profile", encoding="utf-8")

    assert [p.name for p in generate_labels._discover_profiles(tmp_path)] == [
        "first.yaml", "second.yml"
    ]


def test_profile_selection_rejects_profiles_that_do_not_match_sheet(tmp_path):
    import generate_labels

    from shutil import copyfile
    copyfile(PROFILE, tmp_path / "example.yaml")
    book = load_workbook(WORKBOOK)
    book.create_sheet("Unrelated")
    workbook_path = tmp_path / "multi.xlsx"
    book.save(workbook_path)

    try:
        generate_labels._select_profile(tmp_path, workbook_path, "Unrelated")
    except ValueError as exc:
        assert 'No profiles' in str(exc)
        assert 'Unrelated' in str(exc)
    else:
        raise AssertionError("An incompatible profile should not be selectable")
