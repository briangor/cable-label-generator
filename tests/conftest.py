from pathlib import Path
import sys

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def make_profile():
    return {
        "workbook": {"sheet": "Cables"},
        "sections": {
            "device_column": "A",
            "header_match": r"^DEV-",
            "header_row_offset": 1,
            "header_value": "Local Port",
            "data_start_offset": 2,
        },
        "sides": [
            {"name": "left", "device_column": "A", "port_column": "A", "label_column": "B"},
            {"name": "right", "device_column": "I", "port_column": "I", "label_column": "J"},
        ],
        "ordering": {
            "type": "reference",
            "reference_key": "device_port",
            "reference": [["DEV-01", 1], ["DEV-01", 2], ["DEV-02", 1]],
            "unknown_records": "after_reference",
            "unknown_sort": "numeric",
        },
        "layout": {
            "label_group_size": 2,
            "labels_per_table": 4,
            "label_rows": [0, 2],
            "top_label_starts": [1, 4],
            "bottom_label_starts": [0, 3],
            "pages_per_table": 2,
        },
    }


@pytest.fixture
def synthetic_workbook(tmp_path):
    path = tmp_path / "cables.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Cables"

    # Section 1: two left-side and one right-side records.
    ws["A1"] = "DEV-01"
    ws["A2"] = "Local Port"
    ws["A3"] = "25GE1/0/1"
    ws["B3"] = "PEER-A"
    ws["A4"] = "25GE1/0/2"
    ws["B4"] = "PEER-B"
    ws["I3"] = "25GE1/0/10"
    ws["J3"] = "PEER-R1"

    # Section 2: one left-side and one right-side record.
    ws["A6"] = "DEV-02"
    ws["A7"] = "Local Port"
    ws["A8"] = "25GE1/0/1"
    ws["B8"] = "PEER-C"
    ws["I8"] = "25GE1/0/11"
    ws["J8"] = "PEER-R2"

    workbook.save(path)
    return path
