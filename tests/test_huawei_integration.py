from pathlib import Path

import pytest
from docx import Document

import generate_labels


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Kenya_SuperAPP_Phase2_LLD_0901.xlsx"
TEMPLATE = ROOT / "Labels.docx"
PROFILE = ROOT / "profiles" / "huawei_superapp.yaml"

pytestmark = pytest.mark.integration


@pytest.mark.skipif(
    not WORKBOOK.exists() or not TEMPLATE.exists(),
    reason="Private Huawei workbook/template not available",
)
def test_current_huawei_dataset_regression(tmp_path):
    profile = generate_labels.load_profile(PROFILE)
    output = tmp_path / "labels.docx"

    cable_count, label_count, report = generate_labels.generate_document(
        WORKBOOK,
        output,
        TEMPLATE,
        profile,
    )

    assert cable_count == 730
    assert label_count == 1460
    assert len(report["retained"]) == 718
    assert len(report["removed"]) == 6
    assert len(report["added"]) == 12

    document = Document(output)
    assert len(document.tables) == 49
    assert sum(
        1
        for paragraph in document.paragraphs
        if paragraph._p.xpath(".//w:br[@w:type='page']")
    ) == 48
